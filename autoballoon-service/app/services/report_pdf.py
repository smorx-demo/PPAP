"""Generate FAI-style Inspection Report PDF (landscape) — colorful NP1213M-style layout."""

from __future__ import annotations

import io
import re
from pathlib import Path

import fitz

from app.core.config import STATIC

APP_REPORT_TITLE = "AutoBalloon & Inspection Software System"
# Quality certificate / template number printed top-right on every report format
REPORT_CERT_NO = "HQA-TMP12-v3.1"
C_CERT = (0.85, 0.85, 0.85)

# PDF colors (RGB 0–1)
C_LABEL = (0.72, 0.85, 0.95)
C_VALUE = (1.0, 1.0, 1.0)
C_HEADER = (1.0, 0.72, 0.45)
C_ROW_A = (0.82, 0.91, 0.98)
C_ROW_B = (1.0, 1.0, 1.0)
C_BORDER = (0.25, 0.25, 0.25)
C_TEXT = (0, 0, 0)
C_LABEL_TEXT = (0.1, 0.15, 0.25)
C_VERSION_ROW = (1.0, 0.37, 0.08)  # fluorescent orange — Changed / New rows (same as the balloon)
C_VERSION_TEXT = (0.043, 0.18, 0.42)  # dark blue "Version N" (same as the drawing stamp)


def _parse_num(s: str) -> float | None:
    if s is None:
        return None
    t = str(s).strip().replace(",", "")
    if not t:
        return None
    m = re.search(r"[-+]?\d*\.?\d+", t)
    if not m:
        return None
    try:
        return float(m.group(0))
    except ValueError:
        return None


def measure_status(nominal: str, tol_low: str, tol_high: str, measured: str) -> str:
    """
    green = exact match to nominal
    yellow = within tolerance band
    red = outside tolerance / wrong
    empty = no measured value
    """
    m = _parse_num(measured)
    if m is None:
        return "empty"
    n = _parse_num(nominal)
    if n is None:
        return "empty"
    lo = _parse_num(tol_low)
    hi = _parse_num(tol_high)
    if abs(m - n) < 1e-9:
        return "green"
    if lo is not None or hi is not None:
        low = n + (lo if lo is not None else 0.0)
        high = n + (hi if hi is not None else 0.0)
        if lo is not None and hi is not None and lo > hi:
            low, high = n + hi, n + lo
        if lo is None and hi is not None:
            low, high = n - abs(hi), n + abs(hi)
        if hi is None and lo is not None:
            low, high = n - abs(lo), n + abs(lo)
        if low <= m <= high:
            return "yellow"
        return "red"
    return "red"


def _status_color(status: str) -> tuple[float, float, float]:
    return {
        "green": (0.75, 0.95, 0.75),
        "yellow": (1.0, 0.95, 0.65),
        "red": (1.0, 0.78, 0.78),
        "empty": (1, 1, 1),
    }.get(status, (1, 1, 1))


def _cell_text(text) -> str:
    """Visible cell text — blank when empty (never em-dash; Helvetica maps — to ?)."""
    t = str(text or "").strip()
    if not t or t in {"—", "–", "-", "?", "？", "−"}:
        return ""
    # Keep PDF cells short so layout stays fast on huge reports
    if len(t) > 48:
        return t[:45] + "…"
    return t


def _cell(page: fitz.Page, rect: fitz.Rect, fill: tuple[float, float, float], text: str, *, bold_label: bool = False):
    shape = page.new_shape()
    shape.draw_rect(rect)
    shape.finish(color=C_BORDER, fill=fill, width=0.45)
    shape.commit()
    display = _cell_text(text)
    if not display:
        return
    page.insert_textbox(
        rect + (3, 4, -3, -3),
        display,
        fontsize=7 if not bold_label else 7.5,
        fontname="hebo" if bold_label else "helv",
        color=C_TEXT if not bold_label else C_LABEL_TEXT,
        align=0 if bold_label else 1,
    )


def _draw_title_grid(page: fitz.Page, payload: dict, margin: float, page_w: float, y0: float) -> float:
    """Title block grid including Date, Serial Number, and Authorized Signatory."""
    label_w = 92
    val_w = (page_w - 2 * margin - label_w * 2 - 12) / 2
    row_h = 22
    gap = 12
    fields = [
        [("Part Number", payload.get("part_number", "")), ("Part Name", payload.get("part_name", ""))],
        [("Revision", payload.get("revision", "")), ("Material", payload.get("material", ""))],
        [("Mass", payload.get("mass", "")), ("Finish Treatment", payload.get("finish", ""))],
        [("Date", payload.get("report_date", "")), ("Serial Number", payload.get("serial_number", ""))],
        [("Company Name", payload.get("company_name", "")), ("Supplier Name", payload.get("supplier_name", ""))],
        [("Customer Name", payload.get("customer_name", "")), ("Authorized Signatory", payload.get("authorized_signatory", ""))],
    ]
    y = y0
    for row in fields:
        x = margin
        for label, value in row:
            if not label:
                x += label_w + val_w + gap
                continue
            _cell(page, fitz.Rect(x, y, x + label_w, y + row_h), C_LABEL, label, bold_label=True)
            _cell(page, fitz.Rect(x + label_w, y, x + label_w + val_w, y + row_h), C_VALUE, str(value or ""))
            x += label_w + val_w + gap
        y += row_h

    sig_x0 = margin + label_w + val_w + gap + label_w
    sig_x1 = margin + label_w + val_w + gap + label_w + val_w
    sig_y = y + 10
    shape = page.new_shape()
    shape.draw_line(fitz.Point(sig_x0, sig_y), fitz.Point(sig_x1, sig_y))
    shape.finish(color=(0.25, 0.35, 0.40), width=0.8)
    shape.commit()
    page.insert_textbox(
        fitz.Rect(sig_x0, sig_y + 2, sig_x1, sig_y + 16),
        "Authorized Signatory",
        fontsize=8,
        fontname="helv",
        color=(0.25, 0.35, 0.40),
        align=1,
    )
    return (sig_y + 18) - y0 + 8


def _measure_cols(payload: dict) -> list[dict]:
    cols = list(payload.get("measure_columns") or [])
    if cols:
        return cols
    return [
        {"id": "m1", "label": "M1", "values": {}},
        {"id": "m2", "label": "M2", "values": {}},
    ]


def build_report_pdf(payload: dict, logo_path: Path | None = None) -> bytes:
    """
    Build inspection report PDF.
    Optimized for large jobs (thousands of rows): batch cell fills per page
    and avoid one Shape.commit() per cell.
    """
    logo_path = logo_path or (STATIC / "smorx-logo-full.png")
    page_w, page_h = 841.89, 595.28  # A4 landscape
    doc = fitz.open()
    rows = payload.get("rows") or []
    measure_cols = _measure_cols(payload)
    # Pre-resolve measure value maps once (avoid repeated dict lookups)
    measure_maps = [dict(col.get("values") or {}) for col in measure_cols]
    # Drawing version: Changed / New balloon rows are filled fluorescent orange
    version_hl = payload.get("version_highlight") or {}
    version_rows = {str(n) for n in version_hl.get("rows") or []}

    headers = [
        "S.No", "Balloon\nNo.", "Reference", "Nominal", "Tol\n(low)", "Tol\n(high)",
        "Instrument", "Unit", "Instr. ID",
    ]
    for col in measure_cols:
        headers.append(str(col.get("label") or "M"))
    headers.append("Remarks")

    margin = 24
    usable = page_w - 2 * margin
    base_weights = [28, 36, 70, 48, 36, 36, 70, 32, 55] + [52] * len(measure_cols) + [70]
    weights = base_weights[: len(headers)]
    while len(weights) < len(headers):
        weights.insert(-1, 52)
    total_w = sum(weights)
    col_w = [usable * w / total_w for w in weights]

    row_h = 17
    header_h = 30
    top_block_h = 190
    rows_per_page = max(1, int((page_h - margin - top_block_h - header_h - 24) // row_h))

    def draw_page_header(page: fitz.Page, first: bool) -> float:
        y = margin
        cert = fitz.Rect(page_w - margin - 110, y, page_w - margin, y + 16)
        shape = page.new_shape()
        shape.draw_rect(cert)
        shape.finish(color=(0.45, 0.45, 0.45), fill=C_CERT, width=0.5)
        shape.commit()
        if version_hl.get("version"):
            vlabel = f"Version {version_hl['version']}"
            vr = fitz.Rect(cert.x0, cert.y1 + 3, cert.x1, cert.y1 + 19)
            page.draw_rect(vr, color=C_VERSION_TEXT, fill=(1, 1, 1), width=1.2)
            vw = fitz.get_text_length(vlabel, fontname="hebo", fontsize=9)
            page.insert_text(
                fitz.Point(vr.x0 + (vr.width - vw) / 2, vr.y1 - 4.5),
                vlabel,
                fontsize=9,
                fontname="hebo",
                color=C_VERSION_TEXT,
            )
        tw = fitz.get_text_length(REPORT_CERT_NO, fontname="helv", fontsize=8.5)
        page.insert_text(
            fitz.Point(cert.x0 + (cert.width - tw) / 2, cert.y1 - 4.5),
            REPORT_CERT_NO,
            fontsize=8.5,
            fontname="helv",
            color=C_TEXT,
        )
        if first and logo_path.exists():
            try:
                page.insert_image(fitz.Rect(margin, y, margin + 150, y + 38), filename=str(logo_path))
            except Exception:
                pass
        page.insert_text(
            fitz.Point(page_w / 2 - 115, y + 22),
            APP_REPORT_TITLE,
            fontsize=11,
            fontname="hebo",
            color=C_TEXT,
        )
        page.insert_text(
            fitz.Point(page_w / 2 - 42, y + 36),
            "Inspection Report",
            fontsize=9,
            fontname="helv",
            color=(0.35, 0.35, 0.35),
        )
        if not first:
            return 52
        grid_h = _draw_title_grid(page, payload, margin, page_w, y + 44)
        return 44 + grid_h + 8

    measure_start = 9
    measure_end = 9 + len(measure_cols)
    n_rows = max(len(rows), 1)

    for start in range(0, n_rows, rows_per_page):
        page = doc.new_page(width=page_w, height=page_h)
        info_h = draw_page_header(page, first=(start == 0))
        y = margin + info_h
        x0 = margin

        # Header row — one shape for all fills
        shape = page.new_shape()
        x = x0
        header_rects: list[fitz.Rect] = []
        for i, h in enumerate(headers):
            r = fitz.Rect(x, y, x + col_w[i], y + header_h)
            header_rects.append(r)
            shape.draw_rect(r)
            x += col_w[i]
        shape.finish(color=C_BORDER, fill=C_HEADER, width=0.5)
        shape.commit()
        for i, h in enumerate(headers):
            page.insert_textbox(
                header_rects[i] + (2, 4, -2, -2),
                h,
                fontsize=7,
                fontname="hebo",
                color=C_TEXT,
                align=1,
            )
        y += header_h

        chunk = rows[start : start + rows_per_page] if rows else [{}]
        # Batch fills: group cells by fill color, one finish() per color group
        fill_groups: dict[tuple[float, float, float], list[fitz.Rect]] = {}
        text_jobs: list[tuple[fitz.Rect, str]] = []

        for ri, row in enumerate(chunk):
            measured = list(row.get("measured") or [])
            bkey = str(row.get("balloon", ""))
            measure_vals = []
            statuses = []
            for ci, mmap in enumerate(measure_maps):
                val = mmap.get(bkey, "")
                if not val and ci < len(measured):
                    val = measured[ci] or ""
                measure_vals.append(val)
                statuses.append(
                    measure_status(
                        str(row.get("nominal", "")),
                        str(row.get("tol_low", "")),
                        str(row.get("tol_high", "")),
                        val,
                    )
                )

            cells = [
                str(row.get("sno", start + ri + 1)),
                str(row.get("balloon", "")),
                str(row.get("reference", "")),
                str(row.get("nominal", "")),
                str(row.get("tol_low", "")),
                str(row.get("tol_high", "")),
                str(row.get("instrument", "")),
                str(row.get("unit", "")),
                str(row.get("instr_id", "")),
                *measure_vals,
                str(row.get("remarks", "")),
            ]

            row_fill = C_ROW_A if ri % 2 == 0 else C_ROW_B
            if bkey in version_rows:
                row_fill = C_VERSION_ROW
            x = x0
            for i, cell in enumerate(cells):
                r = fitz.Rect(x, y, x + col_w[i], y + row_h)
                fill = row_fill
                if measure_start <= i < measure_end:
                    st = statuses[i - measure_start]
                    if st != "empty":
                        fill = _status_color(st)
                fill_groups.setdefault(fill, []).append(r)
                display = _cell_text(cell)
                if display:
                    text_jobs.append((r, display))
                x += col_w[i]
            y += row_h

        for fill, rects in fill_groups.items():
            shape = page.new_shape()
            for r in rects:
                shape.draw_rect(r)
            shape.finish(color=C_BORDER, fill=fill, width=0.35)
            shape.commit()

        for r, display in text_jobs:
            # insert_text is faster than textbox for short centered-ish values
            page.insert_text(
                fitz.Point(r.x0 + 2, r.y0 + 12),
                display,
                fontsize=7,
                fontname="helv",
                color=C_TEXT,
            )

    buf = io.BytesIO()
    # garbage=1 is much faster than 3 on 80+ page reports
    doc.save(buf, deflate=True, garbage=1)
    doc.close()
    return buf.getvalue()


def _xl_hex(rgb: tuple[float, float, float]) -> str:
    return "".join(f"{round(c * 255):02X}" for c in rgb)


def build_report_xlsx(payload: dict, logo_path: Path | None = None) -> bytes:
    """Excel version of the inspection report with the same layout and colours as the PDF."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    logo_path = logo_path or (STATIC / "smorx-logo-full.png")
    rows = payload.get("rows") or []
    measure_cols = _measure_cols(payload)
    measure_maps = [dict(col.get("values") or {}) for col in measure_cols]

    headers = [
        "S.No", "Balloon\nNo.", "Reference", "Nominal", "Tol\n(low)", "Tol\n(high)",
        "Instrument", "Unit", "Instr. ID",
    ]
    headers += [str(col.get("label") or "M") for col in measure_cols]
    headers.append("Remarks")
    n_cols = len(headers)
    widths = [7, 9, 16, 12, 9, 9, 16, 8, 13] + [12] * len(measure_cols) + [22]

    def fill(rgb: tuple[float, float, float]) -> PatternFill:
        return PatternFill("solid", fgColor=_xl_hex(rgb))

    thin = Side(style="thin", color=_xl_hex(C_BORDER))
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    left = Alignment(horizontal="left", vertical="center", wrap_text=True)

    wb = Workbook()
    ws = wb.active
    ws.title = "Inspection Report"
    ws.sheet_view.showGridLines = False
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w

    # Logo + title
    ws.row_dimensions[1].height = 24
    ws.row_dimensions[2].height = 18
    if logo_path.exists():
        try:
            from openpyxl.drawing.image import Image as XLImage

            img = XLImage(str(logo_path))
            ratio = 40 / max(1, img.height)
            img.height = 40
            img.width = int(img.width * ratio)
            ws.add_image(img, "A1")
        except Exception:
            pass
    title_c0 = 4
    ws.merge_cells(start_row=1, start_column=title_c0, end_row=1, end_column=n_cols - 1)
    ws.merge_cells(start_row=2, start_column=title_c0, end_row=2, end_column=n_cols - 1)
    cert = ws.cell(1, n_cols, REPORT_CERT_NO)
    cert.font = Font(size=9)
    cert.fill = fill(C_CERT)
    cert.border = Border(left=Side(style="thin", color="737373"), right=Side(style="thin", color="737373"),
                         top=Side(style="thin", color="737373"), bottom=Side(style="thin", color="737373"))
    cert.alignment = Alignment(horizontal="center", vertical="center")
    version_hl = payload.get("version_highlight") or {}
    version_rows = {str(n) for n in version_hl.get("rows") or []}
    if version_hl.get("version"):
        vc = ws.cell(2, n_cols, f"Version {version_hl['version']}")
        vc.font = Font(bold=True, size=10, color=_xl_hex(C_VERSION_TEXT))
        vside = Side(style="medium", color=_xl_hex(C_VERSION_TEXT))
        vc.border = Border(left=vside, right=vside, top=vside, bottom=vside)
        vc.alignment = Alignment(horizontal="center", vertical="center")
    t = ws.cell(1, title_c0, APP_REPORT_TITLE)
    t.font = Font(bold=True, size=14)
    t.alignment = Alignment(horizontal="center", vertical="center")
    st = ws.cell(2, title_c0, "Inspection Report")
    st.font = Font(size=11, color="595959")
    st.alignment = Alignment(horizontal="center", vertical="center")

    # Title grid: two label/value pairs per row, split across the table width
    half = n_cols // 2
    pair_spans = [(1, 2, 3, half), (half + 1, half + 2, half + 3, n_cols)]
    fields = [
        [("Part Number", "part_number"), ("Part Name", "part_name")],
        [("Revision", "revision"), ("Material", "material")],
        [("Mass", "mass"), ("Finish Treatment", "finish")],
        [("Date", "report_date"), ("Serial Number", "serial_number")],
        [("Company Name", "company_name"), ("Supplier Name", "supplier_name")],
        [("Customer Name", "customer_name"), ("Authorized Signatory", "authorized_signatory")],
    ]
    r = 4
    for pairs in fields:
        ws.row_dimensions[r].height = 20
        for (label, key), (l0, l1, v0, v1) in zip(pairs, pair_spans):
            ws.merge_cells(start_row=r, start_column=l0, end_row=r, end_column=l1)
            ws.merge_cells(start_row=r, start_column=v0, end_row=r, end_column=v1)
            for c in range(l0, v1 + 1):
                cell = ws.cell(r, c)
                cell.border = border
                cell.fill = fill(C_LABEL if c <= l1 else C_VALUE)
            lc = ws.cell(r, l0, label)
            lc.font = Font(bold=True, color=_xl_hex(C_LABEL_TEXT))
            lc.alignment = left
            vc = ws.cell(r, v0, _cell_text(payload.get(key, "")))
            vc.alignment = center
        r += 1

    # Signature line under the right-hand value column
    r += 1
    sig_c0, sig_c1 = pair_spans[1][2], pair_spans[1][3]
    ws.row_dimensions[r].height = 26
    r += 1
    ws.merge_cells(start_row=r, start_column=sig_c0, end_row=r, end_column=sig_c1)
    sig_line = Side(style="medium", color="40595F")
    for c in range(sig_c0, sig_c1 + 1):
        ws.cell(r, c).border = Border(top=sig_line)
    sc = ws.cell(r, sig_c0, "Authorized Signatory")
    sc.font = Font(size=9, color="40595F")
    sc.alignment = Alignment(horizontal="center", vertical="top")
    r += 2

    # Table header
    header_row = r
    ws.row_dimensions[r].height = 32
    for i, h in enumerate(headers, start=1):
        cell = ws.cell(r, i, h)
        cell.font = Font(bold=True)
        cell.fill = fill(C_HEADER)
        cell.border = border
        cell.alignment = center
    r += 1

    measure_start = 9
    measure_end = 9 + len(measure_cols)
    status_fills = {s: fill(_status_color(s)) for s in ("green", "yellow", "red")}
    band_fills = [fill(C_ROW_A), fill(C_ROW_B)]
    for ri, row in enumerate(rows):
        measured = list(row.get("measured") or [])
        bkey = str(row.get("balloon", ""))
        measure_vals = []
        for ci, mmap in enumerate(measure_maps):
            val = mmap.get(bkey, "")
            if not val and ci < len(measured):
                val = measured[ci] or ""
            measure_vals.append(val)
        cells = [
            row.get("sno", ri + 1),
            row.get("balloon", ""),
            row.get("reference", ""),
            row.get("nominal", ""),
            row.get("tol_low", ""),
            row.get("tol_high", ""),
            row.get("instrument", ""),
            row.get("unit", ""),
            row.get("instr_id", ""),
            *measure_vals,
            row.get("remarks", ""),
        ]
        band = band_fills[ri % 2]
        if bkey in version_rows:
            band = fill(C_VERSION_ROW)
        for i, val in enumerate(cells):
            value = val if isinstance(val, int) else _cell_text(val)
            cell = ws.cell(r, i + 1, value)
            cell.border = border
            cell.alignment = center
            cell.fill = band
            if measure_start <= i < measure_end:
                status = measure_status(
                    str(row.get("nominal", "")),
                    str(row.get("tol_low", "")),
                    str(row.get("tol_high", "")),
                    str(val or ""),
                )
                if status in status_fills:
                    cell.fill = status_fills[status]
        ws.row_dimensions[r].height = 18
        r += 1

    last_col = get_column_letter(n_cols)
    ws.freeze_panes = ws.cell(header_row + 1, 1)
    if rows:
        ws.auto_filter.ref = f"A{header_row}:{last_col}{r - 1}"

    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_title_rows = f"{header_row}:{header_row}"
    ws.print_options.horizontalCentered = True
    ws.page_margins.left = ws.page_margins.right = 0.35
    ws.page_margins.top = ws.page_margins.bottom = 0.5

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def characteristics_to_report_rows(chars: list[dict], zone_map: dict[int, str] | None = None) -> list[dict]:
    rows = []
    zone_map = zone_map or {}
    for i, c in enumerate(chars, start=1):
        text = str(c.get("text") or "")
        nominal = text
        tol_low = ""
        tol_high = ""
        m = re.search(r"([-+]?\d+(?:[.,]\d+)?)\s*±\s*([-+]?\d+(?:[.,]\d+)?)", text)
        if m:
            nominal = m.group(1).replace(",", ".")
            t = m.group(2).replace(",", ".")
            tol_low = f"-{t}"
            tol_high = t
        else:
            m2 = re.search(r"[-+]?\d+(?:[.,]\d+)?", text)
            if m2:
                nominal = m2.group(0).replace(",", ".")
        bno = int(c.get("balloon") or i)
        ref = zone_map.get(bno) or f"P{c.get('page', 1)}"
        rows.append({
            "sno": i,
            "balloon": bno,
            "page": int(c.get("page") or 1),
            "reference": ref,
            "nominal": "" if not text or re.match(r"(?i)^auto\s*balloo?n$", text.strip()) else nominal,
            "tol_low": tol_low,
            "tol_high": tol_high,
            "instrument": "",
            "unit": "mm",
            "instr_id": "",
            "measured": ["", ""],
            "remarks": "",
        })
    return rows


def default_measure_columns() -> list[dict]:
    return [
        {"id": "m1", "label": "M1", "values": {}},
        {"id": "m2", "label": "M2", "values": {}},
    ]


def filter_report_rows_by_page(rows: list[dict], page: int | None) -> list[dict]:
    if not page or page < 1:
        return list(rows)
    filtered = [r for r in rows if int(r.get("page") or 0) == int(page)]
    out = []
    for i, row in enumerate(filtered, start=1):
        nr = dict(row)
        nr["sno"] = i
        out.append(nr)
    return out


def _balloon_in_area_rect(
    b: dict,
    *,
    page: int,
    x: float,
    y: float,
    width: float,
    height: float,
) -> bool:
    if int(b.get("page") or 1) != int(page):
        return False
    cx = float(b.get("x") or 0) + float(b.get("width") or 0) / 2.0
    cy = float(b.get("y") or 0) + float(b.get("height") or 0) / 2.0
    return x <= cx <= x + width and y <= cy <= y + height


def filter_report_rows_by_area(
    rows: list[dict],
    balloons: list[dict],
    *,
    page: int,
    x: float,
    y: float,
    width: float,
    height: float,
) -> list[dict]:
    """Keep report rows whose balloon callout centre lies inside the PDF rectangle."""
    nos = {
        int(b.get("balloon") or 0)
        for b in balloons
        if _balloon_in_area_rect(b, page=page, x=x, y=y, width=width, height=height)
    }
    filtered = [r for r in rows if int(r.get("balloon") or 0) in nos]
    out = []
    for i, row in enumerate(filtered, start=1):
        nr = dict(row)
        nr["sno"] = i
        out.append(nr)
    return out


def filter_report_rows(
    rows: list[dict],
    balloons: list[dict] | None = None,
    *,
    page: int | None = None,
    area_page: int | None = None,
    area_x: float | None = None,
    area_y: float | None = None,
    area_w: float | None = None,
    area_h: float | None = None,
) -> list[dict]:
    """Apply page and/or work-area filters (renumber S.No after each step)."""
    out = list(rows)
    if page and page >= 1:
        out = filter_report_rows_by_page(out, page)
    if (
        area_page
        and area_page >= 1
        and area_x is not None
        and area_y is not None
        and area_w is not None
        and area_h is not None
        and area_w > 0
        and area_h > 0
        and balloons
    ):
        out = filter_report_rows_by_area(
            out,
            balloons,
            page=int(area_page),
            x=float(area_x),
            y=float(area_y),
            width=float(area_w),
            height=float(area_h),
        )
    return out


def report_rows_to_csv(rows: list[dict], measure_columns: list[dict] | None = None) -> bytes:
    import csv
    import io

    measure_columns = measure_columns or []
    headers = [
        "S.No", "Page", "Balloon", "Reference", "Nominal", "Tol low", "Tol high",
        "Instrument", "Unit", "Instr. ID", "Measured 1", "Measured 2", "Remarks",
    ]
    for col in measure_columns:
        headers.insert(-1, col.get("label") or "Measured")
    buf = io.StringIO()
    writer = csv.writer(buf)
    writer.writerow(headers)
    for row in rows:
        measured = list(row.get("measured") or ["", ""])
        while len(measured) < 2:
            measured.append("")
        extra = []
        bkey = str(row.get("balloon"))
        for col in measure_columns:
            extra.append((col.get("values") or {}).get(bkey, ""))
        writer.writerow([
            row.get("sno"), row.get("page", 1), row.get("balloon"), row.get("reference"),
            row.get("nominal"), row.get("tol_low"), row.get("tol_high"),
            row.get("instrument"), row.get("unit"), row.get("instr_id"),
            measured[0], measured[1] if len(measured) > 1 else "", *extra, row.get("remarks"),
        ])
    return buf.getvalue().encode("utf-8-sig")


def report_payload_to_json(
    meta: dict,
    *,
    page: int | None = None,
    area_page: int | None = None,
    area_x: float | None = None,
    area_y: float | None = None,
    area_w: float | None = None,
    area_h: float | None = None,
) -> bytes:
    import json

    rows = filter_report_rows(
        meta.get("report_rows") or [],
        meta.get("balloons") or [],
        page=page,
        area_page=area_page,
        area_x=area_x,
        area_y=area_y,
        area_w=area_w,
        area_h=area_h,
    )
    payload = {
        "part_number": meta.get("part_number", ""),
        "part_name": meta.get("part_name", ""),
        "revision": meta.get("revision", ""),
        "material": meta.get("material", ""),
        "mass": meta.get("mass", ""),
        "finish": meta.get("finish", ""),
        "serial_number": meta.get("serial_number", ""),
        "report_date": meta.get("report_date", ""),
        "company_name": meta.get("company_name", ""),
        "supplier_name": meta.get("supplier_name", ""),
        "customer_name": meta.get("customer_name", ""),
        "authorized_signatory": meta.get("authorized_signatory", ""),
        "balloon_numbering": meta.get("balloon_numbering", "continuous"),
        "page_filter": page,
        "area_filter": (
            {
                "page": area_page,
                "x": area_x,
                "y": area_y,
                "width": area_w,
                "height": area_h,
            }
            if area_page and area_w and area_h
            else None
        ),
        "rows": rows,
        "measure_columns": meta.get("measure_columns") or default_measure_columns(),
        "balloons": meta.get("balloons") or [],
    }
    return json.dumps(payload, indent=2, ensure_ascii=False).encode("utf-8")
