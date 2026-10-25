"""IE Master Sheet inspection report ("INSPECTION REPORT - Mechanical") as PDF and Excel.

The layout is described once as a grid of merged cells per sheet and rendered by both
the PDF (PyMuPDF) and Excel (openpyxl) writers so the two downloads look identical.
"""

from __future__ import annotations

import io
import re
from dataclasses import dataclass

import fitz

from app.services.report_pdf import C_CERT, REPORT_CERT_NO, _cell_text, _measure_cols, measure_status

MASTER_TITLE = "INSPECTION REPORT - Mechanical"
ROWS_PER_SHEET = 12
MIN_ACTUAL_COLS = 5

# Header / footer fields that only exist on the master sheet: (key, label, default)
MASTER_FIELDS: list[tuple[str, str, str]] = [
    ("location", "Location", "ECity/Vemagal"),
    ("supplier", "Supplier Name & Code", "INGENIOUS ENGINEERING PVT LTD"),
    ("qty_offered", "Qty. Offered", "01 No."),
    ("lot_size", "Lot size", "01 No."),
    ("ir_no", "Inspection Report No.", "IE-IR-01"),
    ("sample_size", "Sample size", "01 No."),
    ("prev_fai", "Previous FAI No & Date", "N/A"),
    ("po_no", "PO No", ""),
    ("project_no", "Project No.", ""),
    ("wo_no", "W.O/ Receipt No.", ""),
    ("heat_no", "Heat No.", "N/A"),
    ("route_card", "Route Card No", "N/A"),
    ("test_reports", "Test Reports", ""),
    ("dcr_no", "DCR No", ""),
    ("concession_no", "Concession No", ""),
    ("visual_inspection", "Visual Inspection", ""),
    ("inspected_qty", "Inspected Quantity", "01 No."),
    ("accepted", "Accepted", ""),
    ("rejected", "Rejected", ""),
    ("remarks", "Remarks", ""),
    ("note", "Note", "Please capture all Notes from the drawing"),
    ("inspected_by", "Inspected By", ""),
    ("approved_by", "Witnessed/Approved By", ""),
    ("customer_sign", "Customer (if Applicable)", ""),
    # CMQ_LAB report (same store, cmq_ prefix)
    ("cmq_company", "Company", "Ingenious Engineering Pvt Ltd"),
    ("cmq_doc_no", "Doc No", ""),
    ("cmq_title", "Report title", "CMM Inspection Report"),
    ("cmq_drawing_no", "Drawing Number", ""),
    ("cmq_qty", "Qty", "01 No."),
    ("cmq_inspected_by", "Inspected By", ""),
    ("cmq_approved_by", "Approved By", ""),
    ("cmq_comments", "Comments", ""),
]


def default_master_fields() -> dict[str, str]:
    return {key: default for key, _, default in MASTER_FIELDS}


def ensure_master_fields(meta: dict) -> None:
    """Fill missing master-sheet keys with template defaults (keeps user edits)."""
    current = dict(meta.get("master_fields") or {})
    for key, _, default in MASTER_FIELDS:
        current.setdefault(key, default)
    meta["master_fields"] = current


# ---------------------------------------------------------------------------
# Row content
# ---------------------------------------------------------------------------

def _parameter_name(spec: str) -> str:
    t = (spec or "").strip().upper()
    if not t:
        return ""
    if "Ø" in t or "⌀" in t or "DIA" in t or re.match(r"^(\d+\s*X\s*)?O\s*\d", t):
        return "Diameter"
    if "THK" in t:
        return "Thickness"
    if re.search(r"X\s*45", t):
        return "Chamfer"
    if re.match(r"^(\d+\s*X\s*)?M\d", t):
        return "Thread"
    if re.match(r"^(\d+\s*X\s*)?S?R\s*\d", t):
        return "Radius"
    if "°" in t:
        return "Angle"
    if re.match(r"^RA\b|^RZ\b", t):
        return "Surface Finish"
    if re.search(r"\d", t):
        return "Linear"
    return "Note"


def _signed(v: str) -> str:
    v = v.strip()
    return v if not v or v[0] in "+-" else f"+{v}"


def _tolerance_text(tol_low: str, tol_high: str) -> str:
    lo = str(tol_low or "").strip()
    hi = str(tol_high or "").strip()
    if not lo and not hi:
        return ""
    if lo and hi and lo.lstrip("+-") == hi.lstrip("+-") and lo.startswith("-") and not hi.startswith("-"):
        return f"±{hi.lstrip('+')}"
    return " / ".join(p for p in (_signed(hi), _signed(lo)) if p)


def _spec_text(text: str) -> str:
    return (text or "").replace("⌀", "Ø").strip()


@dataclass
class MasterRow:
    cells: list[str]
    statuses: list[str]  # per actual-dimension column


def build_master_rows(payload: dict, n_actual: int) -> list[MasterRow]:
    rows = payload.get("rows") or []
    measure_cols = _measure_cols(payload)
    measure_maps = [dict(col.get("values") or {}) for col in measure_cols]
    spec_by_balloon: dict = payload.get("spec_by_balloon") or {}
    out: list[MasterRow] = []
    for row in rows:
        bno = row.get("balloon", "")
        key = f"{int(row.get('page') or 1)}:{bno}"
        spec = _spec_text(spec_by_balloon.get(key) or spec_by_balloon.get(str(bno)) or "") or str(
            row.get("nominal", "")
        )
        measured = list(row.get("measured") or [])
        vals: list[str] = []
        statuses: list[str] = []
        for ci in range(n_actual):
            val = ""
            if ci < len(measure_maps):
                val = measure_maps[ci].get(str(bno), "")
                if not val and ci < len(measured):
                    val = measured[ci] or ""
            vals.append(str(val or ""))
            statuses.append(
                measure_status(
                    str(row.get("nominal", "")),
                    str(row.get("tol_low", "")),
                    str(row.get("tol_high", "")),
                    str(val or ""),
                )
            )
        cells = [
            str(bno),
            _parameter_name(spec),
            spec,
            str(row.get("reference", "")),
            _tolerance_text(row.get("tol_low", ""), row.get("tol_high", "")),
            str(row.get("instrument", "")),
            str(row.get("instr_id", "")),
            *vals,
            str(row.get("remarks", "")),
        ]
        out.append(MasterRow(cells=cells, statuses=statuses))
    return out


# ---------------------------------------------------------------------------
# Sheet grid (shared by PDF + Excel)
# ---------------------------------------------------------------------------

@dataclass
class GridCell:
    row: int
    col0: int
    col1: int
    text: str
    style: str  # title | label | value | head | data | box
    rowspan: int = 1
    status: str = "empty"


@dataclass
class SheetGrid:
    col_weights: list[float]
    row_heights: list[float]  # points
    cells: list[GridCell]


def _actual_labels(payload: dict, n_actual: int) -> list[str]:
    labels = [str(c.get("label") or "") for c in _measure_cols(payload)]
    return (labels + [""] * n_actual)[:n_actual]


def build_sheet_grid(
    payload: dict,
    chunk: list[MasterRow],
    sheet_no: int,
    sheet_count: int,
    n_actual: int,
) -> SheetGrid:
    mf = {**default_master_fields(), **(payload.get("master_fields") or {})}
    a0 = 7  # first actual-dimension column
    rem = a0 + n_actual  # remarks column (last)
    weights = [40, 80, 108, 55, 50, 105, 77] + [385 / n_actual] * n_actual + [125]

    cells: list[GridCell] = []
    heights: list[float] = []

    def add(r, c0, c1, text, style, rowspan=1, status="empty"):
        cells.append(GridCell(r, c0, c1, text, style, rowspan, status))

    # Row 0: location + title
    heights.append(24)
    add(0, 0, 1, f"Location :\n{mf['location']}", "label")
    add(0, 2, rem - 1, MASTER_TITLE, "title")
    add(0, rem, rem, REPORT_CERT_NO, "cert")

    # Header block rows 1..8
    heights += [20] * 8
    lot_l, lot_v = a0 + 1, a0 + 2
    ir_l0, ir_l1 = a0 + 3, rem - 1
    heat_v1 = a0 + 2
    sheet_txt = f"{sheet_no:02d} Of {sheet_count:02d}"

    add(1, 0, 1, "Supplier Name & Code:", "label", 2)
    add(1, 2, 4, mf["supplier"], "value", 2)
    add(1, 5, 5, "Qty. Offered:", "label", 2)
    add(1, 6, a0, mf["qty_offered"], "value", 2)
    add(1, lot_l, lot_l, "Lot size:", "label", 2)
    add(1, lot_v, lot_v, mf["lot_size"], "value", 2)
    add(1, ir_l0, ir_l1, "Inspection Report No.:", "label")
    add(1, rem, rem, mf["ir_no"], "value")
    add(2, ir_l0, ir_l1, "Date :", "label")
    add(2, rem, rem, str(payload.get("report_date", "")), "value")

    add(3, 0, 1, "Part/Drg. No :", "label")
    add(3, 2, 4, str(payload.get("part_number", "")), "value")
    add(3, 5, 5, "Issue No.:", "label")
    add(3, 6, a0, str(payload.get("revision", "")), "value")
    add(3, lot_l, lot_l, "Sample size:", "label")
    add(3, lot_v, lot_v, mf["sample_size"], "value")
    add(3, ir_l0, ir_l1, "Previous FAI No & Date", "label")
    add(3, rem, rem, mf["prev_fai"], "value")

    add(4, 0, 1, "Description :", "label")
    add(4, 2, 4, str(payload.get("part_name", "")), "value")
    add(4, 5, 5, "PO No:", "label")
    add(4, 6, a0, mf["po_no"], "value")
    add(4, lot_l, lot_l, "Project No.", "label")
    add(4, lot_v, lot_v, mf["project_no"], "value")
    add(4, ir_l0, ir_l1, "W.O/ Receipt No.:", "label")
    add(4, rem, rem, mf["wo_no"], "value")

    add(5, 0, 1, "Raw Material:", "label")
    add(5, 2, 4, str(payload.get("material", "")), "value")
    add(5, 5, 5, "Heat No.:", "label")
    add(5, 6, heat_v1, mf["heat_no"], "value")
    add(5, ir_l0, ir_l1, "Route Card No:", "label")
    add(5, rem, rem, mf["route_card"], "value")

    add(6, 0, 1, "Finishing", "label", 2)
    add(6, 2, 4, str(payload.get("finish", "")), "value", 2)
    add(6, 5, 5, "Test Reports", "label", 2)
    add(6, 6, heat_v1, mf["test_reports"], "value", 2)
    add(6, ir_l0, ir_l1, "Sheet No.:", "label")
    add(6, rem, rem, sheet_txt, "value")
    add(7, ir_l0, ir_l1, "Part Serial No.:", "label", 2)
    add(7, rem, rem, str(payload.get("serial_number", "")), "value", 2)

    add(8, 0, 1, "DCR No", "label")
    add(8, 2, 4, mf["dcr_no"], "value")
    add(8, 5, 5, "Concession No", "label")
    add(8, 6, heat_v1, mf["concession_no"], "value")

    # Spacer row 9
    heights.append(14)
    add(9, 0, rem, "", "box")

    # Table header rows 10-11
    heights += [20, 16]
    for c, label in enumerate(
        ["SR /\nCH.\nNO.", "Parameters", "Specification", "Zone", "Tol.", "Instrument/\nJig/Fixture", "Inst. ID No"]
    ):
        add(10, c, c, label, "head", 2)
    add(10, a0, rem - 1, "Part SL. No.  &  Actual Dimensions", "head")
    add(10, rem, rem, "Remarks", "head", 2)
    for i, label in enumerate(_actual_labels(payload, n_actual)):
        add(11, a0 + i, a0 + i, label, "head")

    # Data rows
    r = 12
    for i in range(ROWS_PER_SHEET):
        heights.append(16)
        mrow = chunk[i] if i < len(chunk) else None
        for c in range(rem + 1):
            text = mrow.cells[c] if mrow else ""
            status = mrow.statuses[c - a0] if mrow and a0 <= c < rem else "empty"
            add(r, c, c, text, "data", status=status)
        r += 1

    # Footer
    heights.append(30)
    add(r, 0, rem, f"Visual Inspection : {mf['visual_inspection']}", "label")
    r += 1
    heights.append(16)
    add(r, 0, 4, f"Inspected Quantity: {mf['inspected_qty']}", "label")
    add(r, 5, a0, f"Accepted: {mf['accepted']}", "label")
    add(r, a0 + 1, rem, f"Rejected: {mf['rejected']}", "label")
    r += 1
    heights.append(40)
    note = f"Note :{mf['note']}" if mf["note"] else ""
    add(r, 0, rem, f"Remarks: {mf['remarks']}\n\n{note}", "label")
    r += 1
    heights.append(44)
    add(r, 0, 3, f"Inspected By:\n(Signature & Date):\n{mf['inspected_by']}", "sign")
    add(r, 4, a0 + 1, f"Witnessed/Approved By:\n(Signature & Date):\n{mf['approved_by']}", "sign")
    add(r, a0 + 2, rem, f"Customer (if Applicable):\n(Signature & Date):\n{mf['customer_sign']}", "sign")

    return SheetGrid(col_weights=weights, row_heights=heights, cells=cells)


def _sheets(payload: dict) -> tuple[list[SheetGrid], int]:
    n_actual = max(MIN_ACTUAL_COLS, len(_measure_cols(payload)))
    mrows = build_master_rows(payload, n_actual)
    chunks = [mrows[i : i + ROWS_PER_SHEET] for i in range(0, len(mrows), ROWS_PER_SHEET)] or [[]]
    grids = [
        build_sheet_grid(payload, chunk, i + 1, len(chunks), n_actual) for i, chunk in enumerate(chunks)
    ]
    return grids, n_actual


# ---------------------------------------------------------------------------
# Renderers
# ---------------------------------------------------------------------------

STATUS_RGB = {
    "green": (0.75, 0.95, 0.75),
    "yellow": (1.0, 0.95, 0.65),
    "red": (1.0, 0.78, 0.78),
}
TITLE_RGB = (0.94, 0.94, 0.94)
HEAD_RGB = (0.85, 0.85, 0.85)


def _fit_textbox(page: fitz.Page, rect: fitz.Rect, text: str, fontname: str, size: float, align: int) -> None:
    fs = size
    while fs >= 4.5:
        if page.insert_textbox(rect, text, fontsize=fs, fontname=fontname, align=align) >= 0:
            return
        fs -= 0.5
    page.insert_textbox(rect, text[:40] + "…", fontsize=4.5, fontname=fontname, align=align)


def build_master_pdf(payload: dict) -> bytes:
    grids, _ = _sheets(payload)
    return render_pdf(grids, landscape=True)


def render_pdf(grids: list[SheetGrid], *, landscape: bool) -> bytes:
    page_w, page_h = (841.89, 595.28) if landscape else (595.28, 841.89)  # A4
    margin = 18
    doc = fitz.open()
    for grid in grids:
        page = doc.new_page(width=page_w, height=page_h)
        usable_w = page_w - 2 * margin
        total = sum(grid.col_weights)
        xs = [margin]
        for w in grid.col_weights:
            xs.append(xs[-1] + usable_w * w / total)
        scale_h = min(1.0, (page_h - 2 * margin) / sum(grid.row_heights))
        ys = [margin]
        for h in grid.row_heights:
            ys.append(ys[-1] + h * scale_h)

        shape = page.new_shape()
        text_jobs = []
        for cell in grid.cells:
            rect = fitz.Rect(xs[cell.col0], ys[cell.row], xs[cell.col1 + 1], ys[cell.row + cell.rowspan])
            fill = None
            if cell.style == "title":
                fill = TITLE_RGB
            elif cell.style == "cert":
                fill = C_CERT
            elif cell.style == "head":
                fill = HEAD_RGB
            elif cell.style == "data" and cell.status in STATUS_RGB:
                fill = STATUS_RGB[cell.status]
            shape.draw_rect(rect)
            shape.finish(color=(0, 0, 0), fill=fill, width=0.6)
            text = _cell_text(cell.text) if cell.style in ("data", "value") else cell.text
            if text.strip():
                text_jobs.append((rect, text, cell.style))
        shape.commit()

        for rect, text, style in text_jobs:
            inner = rect + (2.5, 2, -2.5, -1)
            if style == "title":
                _fit_textbox(page, rect + (2, 4, -2, -1), text, "hebo", 13, 1)
            elif style == "cert":
                pad = max(1.5, (rect.height - 10) / 2)
                _fit_textbox(page, rect + (2, pad, -2, -1), text, "helv", 8.5, 1)
            elif style == "subtitle":
                pad = max(1.5, (rect.height - 11) / 2)
                _fit_textbox(page, rect + (2, pad, -2, -1), text, "hebo", 9.5, 1)
            elif style == "head":
                lines = text.count("\n") + 1
                pad = max(1.5, (rect.height - lines * 8.5) / 2)
                _fit_textbox(page, rect + (2, pad, -2, -1), text, "hebo", 7.5, 1)
            elif style == "value":
                pad = max(2, (rect.height - 10) / 2)
                _fit_textbox(page, rect + (2.5, pad, -2.5, -1), text, "tibo", 8, 1)
            elif style == "data":
                _fit_textbox(page, rect + (2, 4, -2, -1), text, "helv", 7, 1)
            elif style == "sign":
                _fit_textbox(page, inner, text, "hebo", 8, 0)
            else:
                _fit_textbox(page, inner, text, "hebo", 7, 0)

    buf = io.BytesIO()
    doc.save(buf, deflate=True, garbage=1)
    doc.close()
    return buf.getvalue()


def build_master_xlsx(payload: dict) -> bytes:
    grids, _ = _sheets(payload)
    return render_xlsx(grids, landscape=True)


def render_xlsx(grids: list[SheetGrid], *, landscape: bool) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    def hexrgb(rgb):
        return "".join(f"{round(c * 255):02X}" for c in rgb)

    side = Side(style="thin", color="000000")
    border = Border(left=side, right=side, top=side, bottom=side)
    fills = {
        "title": PatternFill("solid", fgColor=hexrgb(TITLE_RGB)),
        "cert": PatternFill("solid", fgColor=hexrgb(C_CERT)),
        "head": PatternFill("solid", fgColor=hexrgb(HEAD_RGB)),
        **{k: PatternFill("solid", fgColor=hexrgb(v)) for k, v in STATUS_RGB.items()},
    }
    fonts = {
        "title": Font(name="Arial", bold=True, size=14),
        "subtitle": Font(name="Arial", bold=True, size=11),
        "cert": Font(name="Arial", size=9),
        "label": Font(name="Arial", bold=True, size=8),
        "sign": Font(name="Arial", bold=True, size=9),
        "value": Font(name="Times New Roman", bold=True, size=9),
        "head": Font(name="Arial", bold=True, size=9),
        "data": Font(name="Arial", size=9),
        "box": Font(name="Arial", size=8),
    }
    aligns = {
        "title": Alignment(horizontal="center", vertical="center"),
        "subtitle": Alignment(horizontal="center", vertical="center"),
        "cert": Alignment(horizontal="center", vertical="center"),
        "label": Alignment(horizontal="left", vertical="top", wrap_text=True),
        "sign": Alignment(horizontal="left", vertical="top", wrap_text=True),
        "value": Alignment(horizontal="center", vertical="center", wrap_text=True),
        "head": Alignment(horizontal="center", vertical="center", wrap_text=True),
        "data": Alignment(horizontal="center", vertical="center", wrap_text=True),
        "box": Alignment(horizontal="left", vertical="center"),
    }

    sheet_chars = 150 if landscape else 105
    wb = Workbook()
    wb.remove(wb.active)
    for si, grid in enumerate(grids, start=1):
        ws = wb.create_sheet(f"Sheet {si:02d}")
        ws.sheet_view.showGridLines = False
        total = sum(grid.col_weights)
        for i, w in enumerate(grid.col_weights, start=1):
            ws.column_dimensions[get_column_letter(i)].width = round(sheet_chars * w / total, 1)
        for i, h in enumerate(grid.row_heights, start=1):
            ws.row_dimensions[i].height = h
        for cell in grid.cells:
            r0 = cell.row + 1
            r1 = cell.row + cell.rowspan
            c0 = cell.col0 + 1
            c1 = cell.col1 + 1
            if r1 > r0 or c1 > c0:
                ws.merge_cells(start_row=r0, start_column=c0, end_row=r1, end_column=c1)
            fill = fills.get(cell.style)
            if cell.style == "data" and cell.status in STATUS_RGB:
                fill = fills[cell.status]
            for rr in range(r0, r1 + 1):
                for cc in range(c0, c1 + 1):
                    xc = ws.cell(rr, cc)
                    xc.border = border
                    if fill is not None:
                        xc.fill = fill
            text = _cell_text(cell.text) if cell.style in ("data", "value") else cell.text
            tc = ws.cell(r0, c0, text or None)
            tc.font = fonts[cell.style]
            tc.alignment = aligns[cell.style]

        ws.page_setup.orientation = "landscape" if landscape else "portrait"
        ws.page_setup.paperSize = ws.PAPERSIZE_A4
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 1
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.print_options.horizontalCentered = True
        ws.page_margins.left = ws.page_margins.right = 0.25
        ws.page_margins.top = ws.page_margins.bottom = 0.3
        ws.print_area = f"A1:{get_column_letter(len(grid.col_weights))}{len(grid.row_heights)}"

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
